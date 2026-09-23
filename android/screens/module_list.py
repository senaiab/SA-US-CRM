import os

from kivy.app import App
from kivy.lang import Builder
from kivy.metrics import dp
from kivy.properties import StringProperty
from kivy.uix.behaviors import ToggleButtonBehavior
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.screenmanager import Screen

import db as dbmod

KV = """
<RecordRow>:
    size_hint_y: None
    height: dp(56)
    group: "record_rows"
    text: ""
    halign: "left"
    valign: "middle"
    text_size: self.width - dp(16), None
    padding: dp(8), dp(8)

<ModuleListScreen>:
    BoxLayout:
        orientation: "vertical"

        BoxLayout:
            size_hint_y: None
            height: dp(52)
            padding: dp(6)
            canvas.before:
                Color:
                    rgba: 0.122, 0.306, 0.471, 1
                Rectangle:
                    pos: self.pos
                    size: self.size
            Button:
                text: "< Back"
                size_hint_x: None
                width: dp(80)
                on_release: root.go_back()
            Label:
                text: root.display_title
                bold: True

        BoxLayout:
            size_hint_y: None
            height: dp(48)
            padding: dp(6)
            spacing: dp(6)
            TextInput:
                id: search_box
                hint_text: "Search"
                multiline: False
                on_text_validate: root.do_search()
            Button:
                text: "Search"
                size_hint_x: None
                width: dp(80)
                on_release: root.do_search()
            Button:
                text: "Clear"
                size_hint_x: None
                width: dp(70)
                on_release: root.clear_search()

        ScrollView:
            BoxLayout:
                id: record_list
                orientation: "vertical"
                size_hint_y: None
                height: self.minimum_height

        BoxLayout:
            id: toolbar
            size_hint_y: None
            height: dp(52)
            padding: dp(6)
            spacing: dp(6)

        Label:
            id: status_label
            text: ""
            size_hint_y: None
            height: dp(24)
            color: 0.4, 0.4, 0.4, 1
"""
Builder.load_string(KV)


class RecordRow(ToggleButtonBehavior, Label):
    record_id = None


class ModuleListScreen(Screen):
    module = StringProperty("")
    display_title = StringProperty("")

    def __init__(self, module="", **kwargs):
        self.module = module
        self.display_title = dbmod.display_name(module)
        super().__init__(**kwargs)
        self._records = []
        self._built_toolbar = False

    def on_pre_enter(self):
        if not self._built_toolbar:
            self._build_toolbar()
            self._built_toolbar = True

    def _build_toolbar(self):
        toolbar = self.ids.toolbar
        toolbar.clear_widgets()
        toolbar.add_widget(self._btn("Add", self.add_record))
        toolbar.add_widget(self._btn("Edit", self.edit_selected))
        toolbar.add_widget(self._btn("Delete", self.delete_selected))
        if self.module == "clients":
            toolbar.add_widget(self._btn("New Quotation", self.new_quotation_for_client))
            toolbar.add_widget(self._btn("Import Excel", self.import_excel))
        toolbar.add_widget(self._btn("Export Excel", self.export_excel))

    @staticmethod
    def _btn(text, handler):
        b = Button(text=text)
        b.bind(on_release=lambda inst: handler())
        return b

    def refresh(self):
        self.do_search()

    def do_search(self):
        app = App.get_running_app()
        q = self.ids.search_box.text.strip()
        self._records = app.database.list_records(self.module, q)
        self._populate()

    def clear_search(self):
        self.ids.search_box.text = ""
        self.do_search()

    def _populate(self):
        fields = dbmod.MODULES[self.module]
        container = self.ids.record_list
        container.clear_widgets()
        for rec in self._records:
            d = rec.get("data", {})
            summary = " | ".join(str(d.get(f, "")) for f in fields[:3] if d.get(f))
            row = RecordRow(text=summary or f"Record #{rec['id']}")
            row.record_id = rec["id"]
            container.add_widget(row)
        self.ids.status_label.text = f"{len(self._records)} record{'s' if len(self._records) != 1 else ''}"

    def selected_record(self):
        for widget in self.ids.record_list.children:
            if isinstance(widget, RecordRow) and widget.state == "down":
                return next((r for r in self._records if r["id"] == widget.record_id), None)
        return None

    def add_record(self):
        App.get_running_app().open_record_form(self.module)

    def edit_selected(self):
        rec = self.selected_record()
        if not rec:
            self._notify("Select a record first.")
            return
        App.get_running_app().open_record_form(self.module, record=rec)

    def delete_selected(self):
        rec = self.selected_record()
        if not rec:
            self._notify("Select a record first.")
            return
        app = App.get_running_app()

        def confirm(_):
            try:
                app.database.delete_record(rec["id"], app.current_user["username"], app.current_user["role"])
            except (dbmod.AuthError, dbmod.ValidationError) as e:
                popup.dismiss()
                self._notify(str(e))
                return
            popup.dismiss()
            self.refresh()

        box = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(12))
        box.add_widget(Label(text="Delete the selected record?\nOnly Owner/Business Manager can delete."))
        btn_row = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(8))
        yes_btn = Button(text="Delete")
        no_btn = Button(text="Cancel")
        btn_row.add_widget(yes_btn)
        btn_row.add_widget(no_btn)
        box.add_widget(btn_row)
        popup = Popup(title="Confirm delete", content=box, size_hint=(0.8, 0.4))
        yes_btn.bind(on_release=confirm)
        no_btn.bind(on_release=lambda inst: popup.dismiss())
        popup.open()

    def new_quotation_for_client(self):
        rec = self.selected_record()
        if not rec:
            self._notify("Select a client first.")
            return
        name = rec.get("data", {}).get("Client Name", "")
        App.get_running_app().open_record_form("quotations", record=None, prefill={"Client Name": name})

    def export_excel(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self._notify("openpyxl is not available.")
            return
        app = App.get_running_app()
        fields = dbmod.MODULES[self.module]
        records = app.database.list_records(self.module, "")
        wb = Workbook()
        ws = wb.active
        ws.title = self.module
        ws.append(["ID"] + fields)
        for r in records:
            ws.append([r["id"]] + [r.get("data", {}).get(f, "") for f in fields])
        out_dir = app.user_data_dir
        path = os.path.join(out_dir, f"{self.module}_export.xlsx")
        wb.save(path)
        self._notify(f"Exported {len(records)} record(s) to {path}")

    def import_excel(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self._notify("openpyxl is not available.")
            return
        app = App.get_running_app()
        try:
            from plyer import filechooser
            paths = filechooser.open_file(filters=[("Excel files", "*.xlsx")])
        except Exception:
            paths = None
        if not paths:
            fallback = os.path.join(app.user_data_dir, "clients_import.xlsx")
            self._notify(f"No file picker available. Place a file at:\n{fallback}")
            path = fallback
            if not os.path.exists(path):
                return
        else:
            path = paths[0]

        fields = dbmod.MODULES["clients"]
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            self._notify("The spreadsheet is empty.")
            return
        field_by_lower = {f.lower(): f for f in fields}
        col_map = {}
        for idx, h in enumerate(header):
            if h is None:
                continue
            field = field_by_lower.get(str(h).strip().lower())
            if field:
                col_map[idx] = field
        if not col_map:
            self._notify("No matching columns found. Expected headers: " + ", ".join(fields))
            return

        existing = app.database.list_records("clients", "")
        existing_keys = {
            (r.get("data", {}).get("Client Name", "").strip().lower(), r.get("data", {}).get("Company", "").strip().lower())
            for r in existing
        }
        imported = skipped = 0
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            data = {f: "" for f in fields}
            for idx, field in col_map.items():
                if idx < len(row) and row[idx] is not None:
                    data[field] = str(row[idx]).strip()
            name = data.get("Client Name", "").strip()
            if not name:
                continue
            key = (name.lower(), data.get("Company", "").strip().lower())
            if key in existing_keys:
                skipped += 1
                continue
            app.database.create_record("clients", data, app.current_user["username"])
            existing_keys.add(key)
            imported += 1
        self._notify(f"Imported {imported} client(s). Skipped {skipped} duplicate(s).")
        self.refresh()

    def go_back(self):
        App.get_running_app().go_to_dashboard()

    def _notify(self, message):
        box = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(12))
        box.add_widget(Label(text=message))
        btn = Button(text="OK", size_hint_y=None, height=dp(44))
        box.add_widget(btn)
        popup = Popup(title="", content=box, size_hint=(0.85, 0.4))
        btn.bind(on_release=lambda inst: popup.dismiss())
        popup.open()
