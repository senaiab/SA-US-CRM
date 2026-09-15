import configparser
import csv
import json
import os
import sys
import threading
from pathlib import Path
from tkinter import *
from tkinter import ttk, messagebox, filedialog
from datetime import datetime

PREFS_FILE = Path.home() / ".sa_us_erp" / "user_prefs.json"


def load_prefs():
    try:
        if PREFS_FILE.exists():
            return json.loads(PREFS_FILE.read_text(encoding="utf-8"))
    except Exception:
        pass
    return {}


def save_prefs(data):
    try:
        PREFS_FILE.parent.mkdir(parents=True, exist_ok=True)
        PREFS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


import requests
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageTk

if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

config = configparser.ConfigParser()
config.read(BASE_DIR / "client_config.ini")
SERVER_URL = config.get("server", "url", fallback="http://127.0.0.1:8080").rstrip("/")

BRANDING_FILE = BASE_DIR / "tenant_branding.json"
DEFAULT_BRANDING = {
    "software_name": "SA-US ERP",
    "developer_name": "SA Ultimate Solutions",
    "client_company_name": "Your Company Name",
    "company_logo_path": "company_logo.png",
    "app_logo_path": "branding/app_logo.png",
    "developer_logo_path": "branding/sa_ultimate_solutions_logo.png",
    "primary_color": "#1F4E78",
}


def load_branding():
    data = DEFAULT_BRANDING.copy()
    if BRANDING_FILE.exists():
        try:
            loaded = json.loads(BRANDING_FILE.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                data.update(loaded)
        except Exception:
            pass
    return data


BRANDING = load_branding()
SOFTWARE_NAME = BRANDING["software_name"]
DEVELOPER_NAME = BRANDING["developer_name"]
CLIENT_COMPANY_NAME = BRANDING.get("client_company_name", "Your Company Name")
PRIMARY_COLOR = BRANDING.get("primary_color", "#1F4E78")


def resolve_asset(path_text):
    if not path_text:
        return None
    p = Path(path_text)
    if not p.is_absolute():
        p = BASE_DIR / p
    return p if p.exists() else None


APP_LOGO_PATH = resolve_asset(BRANDING.get("app_logo_path"))
DEVELOPER_LOGO_PATH = resolve_asset(BRANDING.get("developer_logo_path"))
COMPANY_LOGO_PATH = resolve_asset(BRANDING.get("company_logo_path"))


def set_window_icon(window):
    try:
        if APP_LOGO_PATH and APP_LOGO_PATH.exists():
            img = Image.open(APP_LOGO_PATH).resize((64, 64))
            photo = ImageTk.PhotoImage(img)
            window.iconphoto(True, photo)
            window._icon_photo = photo
    except Exception:
        pass


def load_tk_image(path_obj, size):
    try:
        if path_obj and path_obj.exists():
            img = Image.open(path_obj).convert("RGBA")
            img.thumbnail(size)
            return ImageTk.PhotoImage(img)
    except Exception:
        pass
    return None


DROPDOWNS = {
    "Source": ["Walk-In", "Phone", "Email", "WhatsApp", "Telegram", "Referral", "Repeat Client"],
    "Service Type": ["Air Ticket", "Hotel", "Tour Package", "Visa Support", "Transfer", "Insurance", "Other"],
    "Priority": ["Low", "Normal", "High", "Urgent", "VIP"],
    "Lead Status": ["New", "Qualified", "Quoted", "Follow-Up", "Confirmed", "Closed-Won", "Closed-Lost", "Dormant"],
    "Booked?": ["Yes", "No"],
    "Outcome": ["Won", "Lost-Price", "Lost-Timing", "Lost-No Response", "Cancelled"],
    "Quote Status": ["Draft", "Sent", "Follow-Up", "Accepted", "Expired", "Lost"],
    "Payment Status": ["Pending", "Part Paid", "Paid", "Credit Approved", "Overdue", "Refund Pending", "Refunded"],
    "Credit Approved?": ["Yes", "No"],
    "Handover Complete?": ["Yes", "No"],
    "Issue Status": ["Pending", "Confirmed", "Issued", "Reissued", "Cancelled", "Refunded"],
    "Refund Status": ["None", "Pending", "Submitted", "Approved", "Paid"],
    "AR Stage": ["Current", "Reminder 1", "Reminder 2", "Escalated", "Suspended"],
    "Severity": ["Level 1", "Level 2", "Level 3", "Legal", "Public"],
    "Status": ["Open", "Investigating", "Pending Client", "Resolved", "Closed", "Escalated"],
    "Client Informed?": ["Yes", "No"],
    "Type": ["Discount", "Credit", "Refund/Waiver", "Policy Override", "Client Dispute", "Other"],
    "Closed?": ["Yes", "No"],
    "Client Type": ["Retail", "Corporate Prepaid", "Corporate Credit", "VIP", "Strategic"],
}

# Row color tags: keyed by tag name → (background, foreground)
ROW_TAGS = {
    "green":  ("#DCFCE7", "#166534"),
    "amber":  ("#FEF9C3", "#854D0E"),
    "red":    ("#FEE2E2", "#991B1B"),
    "muted":  ("#F1F5F9", "#94A3B8"),
}

# Dashboard metric cards → module to jump to on click
CARD_MODULE_MAP = {
    "Open Leads":      "leads",
    "Overdue AR":      "ar",
    "Open Complaints": "complaints",
    "Open Exceptions": "exceptions",
}


def _row_tag(module: str, data: dict) -> str:
    """Return a color tag name for a treeview row based on its status values."""
    if module == "leads":
        status = data.get("Lead Status", "")
        priority = data.get("Priority", "")
        if status == "Closed-Won":
            return "green"
        if status in ("Closed-Lost", "Dormant"):
            return "muted"
        if priority in ("Urgent", "VIP"):
            return "red"
        if priority == "High":
            return "amber"
    elif module == "quotations":
        status = data.get("Quote Status", "")
        if status == "Accepted":
            return "green"
        if status in ("Lost", "Expired"):
            return "muted"
        if status == "Follow-Up":
            return "amber"
    elif module == "bookings":
        pay = data.get("Payment Status", "")
        issue = data.get("Issue Status", "")
        if pay == "Overdue":
            return "red"
        if pay == "Paid" and issue == "Issued":
            return "green"
        if issue == "Cancelled":
            return "muted"
    elif module == "ar":
        pay = data.get("Payment Status", "")
        stage = data.get("AR Stage", "")
        if pay == "Overdue":
            return "red"
        if pay == "Paid":
            return "green"
        if stage in ("Escalated", "Suspended"):
            return "amber"
    elif module == "complaints":
        severity = data.get("Severity", "")
        status = data.get("Status", "")
        if severity in ("Legal", "Public"):
            return "red"
        if status in ("Resolved", "Closed"):
            return "green"
        if severity == "Level 3":
            return "amber"
    elif module == "exceptions":
        if data.get("Closed?") == "Yes":
            return "green"
    return ""


# ─── Exceptions ──────────────────────────────────────────────────────────────

class TokenExpiredError(Exception):
    """Raised when the server returns 401 — session has expired."""
    pass


# ─── API ─────────────────────────────────────────────────────────────────────

class API:
    def __init__(self):
        self.token = None
        self.role = None
        self.full_name = None
        self.must_change_password = False

    def headers(self):
        return {"Authorization": f"Bearer {self.token}"} if self.token else {}

    def _handle(self, r):
        """Check for 401 (expired session) before raising generic HTTP errors."""
        if r.status_code == 401:
            self.token = None
            raise TokenExpiredError("Your session has expired. Please log in again.")
        r.raise_for_status()
        return r

    def check_health(self):
        r = requests.get(f"{SERVER_URL}/health", timeout=5)
        r.raise_for_status()

    def login(self, username, password):
        r = requests.post(f"{SERVER_URL}/token", data={"username": username, "password": password}, timeout=10)
        if r.status_code == 429:
            raise Exception(r.json().get("detail", "Too many login attempts. Wait 5 minutes."))
        if r.status_code != 200:
            raise Exception(r.json().get("detail", "Login failed"))
        data = r.json()
        self.token = data["access_token"]
        self.role = data.get("role")
        self.full_name = data.get("full_name")
        self.must_change_password = bool(data.get("must_change_password", False))
        return data

    def modules(self):
        r = requests.get(f"{SERVER_URL}/modules", headers=self.headers(), timeout=10)
        return self._handle(r).json()["modules"]

    def dashboard(self):
        r = requests.get(f"{SERVER_URL}/dashboard", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def list_records(self, module, q=""):
        r = requests.get(f"{SERVER_URL}/records/{module}", params={"q": q}, headers=self.headers(), timeout=15)
        return self._handle(r).json()

    def delete_record(self, record_id):
        r = requests.delete(f"{SERVER_URL}/records/{record_id}", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def create_user(self, username, full_name, role, password):
        payload = {"username": username, "full_name": full_name, "role": role, "new_password": password}
        r = requests.post(f"{SERVER_URL}/admin/user", json=payload, headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def reset_user_password(self, username, new_password):
        """Admin resets an existing user's password without touching role/full name."""
        payload = {"username": username, "new_password": new_password}
        r = requests.post(f"{SERVER_URL}/admin/user", json=payload, headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def change_own_password(self, current_password, new_password):
        payload = {"current_password": current_password, "new_password": new_password}
        r = requests.put(f"{SERVER_URL}/users/me/password", json=payload, headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def save_record(self, module, fields, data, record_id=None, version=None):
        code = data.get(fields[0], "") if fields else ""
        title = data.get("Client Name") or data.get("Company") or data.get("Supplier Name") or code
        payload = {"module": module, "code": code, "title": title, "data": data}
        if version is not None:
            payload["version"] = version
        if record_id:
            r = requests.put(f"{SERVER_URL}/records/{record_id}", json=payload, headers=self.headers(), timeout=10)
        else:
            r = requests.post(f"{SERVER_URL}/records", json=payload, headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def list_users(self):
        r = requests.get(f"{SERVER_URL}/admin/users", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def deactivate_user(self, username):
        r = requests.post(f"{SERVER_URL}/admin/users/{username}/deactivate", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def activate_user(self, username):
        r = requests.post(f"{SERVER_URL}/admin/users/{username}/activate", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def get_audit(self, module=None, username=None):
        params = {}
        if module:
            params["module"] = module
        if username:
            params["username"] = username
        r = requests.get(f"{SERVER_URL}/admin/audit", params=params, headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def list_trash(self, module):
        r = requests.get(f"{SERVER_URL}/trash/{module}", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def restore_record(self, record_id):
        r = requests.post(f"{SERVER_URL}/records/{record_id}/restore", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def reminders(self):
        r = requests.get(f"{SERVER_URL}/reminders", headers=self.headers(), timeout=10)
        return self._handle(r).json()

    def global_search(self, q):
        r = requests.get(f"{SERVER_URL}/search", params={"q": q}, headers=self.headers(), timeout=15)
        return self._handle(r).json()


# ─── Force Password Change Dialog ────────────────────────────────────────────

class ForcePasswordChangeDialog(Toplevel):
    """Shown when the server flags must_change_password. User cannot dismiss without changing."""
    def __init__(self, login_window, api):
        super().__init__(login_window)
        self.login_window = login_window
        self.api = api
        self.title("Change Password Required")
        self.geometry("420x320")
        self.resizable(False, False)
        self.configure(bg="#f3f6fb")
        self.protocol("WM_DELETE_WINDOW", lambda: None)  # prevent close
        set_window_icon(self)
        self.grab_set()

        Label(self, text="You must set a new password before continuing.", font=("Segoe UI", 10, "bold"),
              bg="#f3f6fb", fg="#dc2626", wraplength=380, justify=LEFT).pack(pady=(20, 10), padx=20)

        frm = Frame(self, bg="#f3f6fb")
        frm.pack(padx=20, fill=X)

        Label(frm, text="Current Password", bg="#f3f6fb").pack(anchor=W)
        self.current = Entry(frm, show="*")
        self.current.pack(fill=X, pady=(0, 8))

        Label(frm, text="New Password", bg="#f3f6fb").pack(anchor=W)
        self.new_pw = Entry(frm, show="*")
        self.new_pw.pack(fill=X, pady=(0, 8))

        Label(frm, text="Confirm New Password", bg="#f3f6fb").pack(anchor=W)
        self.confirm = Entry(frm, show="*")
        self.confirm.pack(fill=X, pady=(0, 12))

        self.btn = Button(frm, text="Set Password", bg=PRIMARY_COLOR, fg="white", command=self._submit)
        self.btn.pack(fill=X, ipady=6)

        self.status = Label(frm, text="", bg="#f3f6fb", fg="#dc2626", font=("Segoe UI", 9))
        self.status.pack(pady=(6, 0))
        self.current.focus_set()

    def _submit(self):
        cur, new, conf = self.current.get(), self.new_pw.get(), self.confirm.get()
        if not new or not cur:
            self.status.config(text="All fields are required.")
            return
        if new != conf:
            self.status.config(text="New passwords do not match.")
            return
        if len(new) < 8:
            self.status.config(text="Password must be at least 8 characters.")
            return
        self.btn.config(state=DISABLED, text="Saving...")
        def attempt():
            try:
                self.api.change_own_password(cur, new)
                self.api.must_change_password = False
                self.after(0, self._done)
            except Exception as e:
                self.after(0, lambda: (self.status.config(text=str(e)), self.btn.config(state=NORMAL, text="Set Password")))
        threading.Thread(target=attempt, daemon=True).start()

    def _done(self):
        self.destroy()
        self.login_window.destroy()
        MainWindow(self.api).mainloop()


# ─── Login Window ─────────────────────────────────────────────────────────────

class LoginWindow(Tk):
    def __init__(self, api):
        super().__init__()
        self.api = api
        self.title(f"{SOFTWARE_NAME} — Login")
        self.geometry("520x640")
        self.resizable(False, False)
        self.configure(bg="#f3f6fb")
        set_window_icon(self)

        self.app_logo_img = load_tk_image(APP_LOGO_PATH, (120, 120))
        self.dev_logo_img = load_tk_image(DEVELOPER_LOGO_PATH, (260, 80))
        self.company_logo_img = load_tk_image(COMPANY_LOGO_PATH, (240, 80))

        if self.app_logo_img:
            Label(self, image=self.app_logo_img, bg="#f3f6fb").pack(pady=(18, 6))

        Label(self, text=SOFTWARE_NAME, font=("Segoe UI", 24, "bold"), fg=PRIMARY_COLOR, bg="#f3f6fb").pack(pady=(4, 2))

        # Developer branding — always shown
        Label(self, text="Developed and distributed by", font=("Segoe UI", 10), fg="#475569", bg="#f3f6fb").pack()
        Label(self, text=DEVELOPER_NAME, font=("Segoe UI", 11, "bold"), fg="#0f172a", bg="#f3f6fb").pack(pady=(0, 6))
        if self.dev_logo_img:
            Label(self, image=self.dev_logo_img, bg="#f3f6fb").pack(pady=(0, 10))

        # Client company branding
        Label(self, text=CLIENT_COMPANY_NAME, font=("Segoe UI", 12, "bold"), fg=PRIMARY_COLOR, bg="#f3f6fb").pack()
        if self.company_logo_img:
            Label(self, image=self.company_logo_img, bg="#f3f6fb").pack(pady=(4, 8))

        Label(self, text=f"Server: {SERVER_URL}", bg="#f3f6fb", fg="#334155", font=("Segoe UI", 9)).pack(pady=(0, 10))

        frm = Frame(self, bg="#f3f6fb")
        frm.pack(padx=35, fill=X)
        Label(frm, text="Username", bg="#f3f6fb").pack(anchor=W)
        self.username = Entry(frm)
        self.username.pack(fill=X, pady=(0, 10))
        Label(frm, text="Password", bg="#f3f6fb").pack(anchor=W)
        self.password = Entry(frm, show="*")
        self.password.pack(fill=X, pady=(0, 12))

        self.login_btn = Button(frm, text="Login", bg=PRIMARY_COLOR, fg="white", command=self.do_login)
        self.login_btn.pack(fill=X, ipady=6)

        self.status_label = Label(frm, text="", bg="#f3f6fb", fg="#dc2626", font=("Segoe UI", 9))
        self.status_label.pack(pady=(8, 0))

        self.bind("<Return>", lambda e: self.do_login())
        self.username.focus_set()

        # Check server reachability on open
        self._check_server()

    def _check_server(self):
        def check():
            try:
                self.api.check_health()
                self.after(0, lambda: self.status_label.config(text="Server reachable", fg="#16a34a"))
            except Exception:
                self.after(0, lambda: self.status_label.config(
                    text=f"Cannot reach server at {SERVER_URL}\nCheck client_config.ini and ensure the server is running.", fg="#dc2626"
                ))
        threading.Thread(target=check, daemon=True).start()

    def do_login(self):
        self.login_btn.config(state=DISABLED, text="Logging in...")
        self.status_label.config(text="", fg="#dc2626")

        def attempt():
            try:
                self.api.login(self.username.get(), self.password.get())
                self.after(0, self._on_login_success)
            except Exception as e:
                self.after(0, lambda: self._on_login_fail(str(e)))

        threading.Thread(target=attempt, daemon=True).start()

    def _on_login_success(self):
        if self.api.must_change_password:
            ForcePasswordChangeDialog(self, self.api)
        else:
            self.destroy()
            MainWindow(self.api).mainloop()

    def _on_login_fail(self, msg):
        self.login_btn.config(state=NORMAL, text="Login")
        self.status_label.config(text=msg, fg="#dc2626")


# ─── Record Dialog ────────────────────────────────────────────────────────────

# Modules where "Client Name" should autocomplete against existing Clients records
CLIENT_LOOKUP_MODULES = ("leads", "quotations")


class RecordDialog(Toplevel):
    def __init__(self, parent, module, fields, api, record=None, prefill=None):
        super().__init__(parent)
        self.main_window = parent
        self.module = module
        self.fields = fields
        self.api = api
        self.record = record
        self.title(("Edit" if record else "Add") + f" {module.title()}")
        self.geometry("720x620")
        self.entries = {}

        canvas = Canvas(self)
        scroll = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        frame = Frame(canvas)
        frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=frame, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side=LEFT, fill=BOTH, expand=True)
        scroll.pack(side=RIGHT, fill=Y)

        data = (record or {}).get("data", {}) or (prefill or {})
        for i, field in enumerate(fields):
            Label(frame, text=field, font=("Segoe UI", 9, "bold")).grid(row=i, column=0, sticky=W, padx=12, pady=6)
            if field == "Client Name" and module in CLIENT_LOOKUP_MODULES:
                client_names = sorted({
                    c.get("data", {}).get("Client Name", "").strip()
                    for c in getattr(self.main_window, "records", {}).get("clients", [])
                    if c.get("data", {}).get("Client Name", "").strip()
                })
                var = StringVar(value=str(data.get(field, "")))
                entry = ttk.Combobox(frame, textvariable=var, values=client_names, width=58)
                entry.bind("<<ComboboxSelected>>", self._autofill_from_client)
                entry.bind("<FocusOut>", self._autofill_from_client)
            elif field in DROPDOWNS:
                var = StringVar(value=str(data.get(field, "")))
                entry = ttk.Combobox(frame, textvariable=var, values=DROPDOWNS[field], width=58)
            elif field in ["Notes", "Remarks", "Reason", "Resolution", "Corrective Action", "Performance Notes", "Billing Address"]:
                entry = Text(frame, width=58, height=3)
                entry.insert("1.0", str(data.get(field, "")))
            else:
                entry = Entry(frame, width=62)
                entry.insert(0, str(data.get(field, "")))
            entry.grid(row=i, column=1, sticky=EW, padx=12, pady=6)
            self.entries[field] = entry

        btn_frame = Frame(frame)
        btn_frame.grid(row=len(fields), column=0, columnspan=2, sticky=EW, padx=12, pady=16)
        self.save_btn = Button(btn_frame, text="Save", bg=PRIMARY_COLOR, fg="white", command=self.save)
        self.save_btn.pack(side=LEFT, padx=4, ipadx=18, ipady=4)
        Button(btn_frame, text="Cancel", command=self.destroy).pack(side=LEFT, padx=4, ipadx=18, ipady=4)

    def _autofill_from_client(self, event=None):
        """When Client Name matches an existing Client record, pull its details into
        empty fields on this form so staff don't retype what's already on file."""
        if self.module not in CLIENT_LOOKUP_MODULES:
            return
        name = self.entries["Client Name"].get().strip()
        if not name:
            return
        clients = getattr(self.main_window, "records", {}).get("clients", [])
        match = next(
            (c for c in clients if c.get("data", {}).get("Client Name", "").strip().lower() == name.lower()),
            None,
        )
        if not match:
            return
        cdata = match.get("data", {})
        if self.module == "leads":
            mapping = {
                "Company": cdata.get("Company", "").strip(),
                "Phone/Email": " / ".join(p for p in [cdata.get("Phone", "").strip(), cdata.get("Email", "").strip()] if p),
            }
            for field, value in mapping.items():
                widget = self.entries.get(field)
                if widget is not None and value and not widget.get().strip():
                    widget.delete(0, END)
                    widget.insert(0, value)

    def save(self):
        data = {}
        for f, widget in self.entries.items():
            if isinstance(widget, Text):
                data[f] = widget.get("1.0", END).strip()
            else:
                data[f] = widget.get().strip()
        self.save_btn.config(state=DISABLED, text="Saving...")

        def do_save():
            try:
                record_id = self.record.get("id") if self.record else None
                version = self.record.get("version") if self.record else None
                self.api.save_record(self.module, self.fields, data, record_id, version=version)
                self.after(0, lambda: (self.main_window.refresh_current_tab(), self.destroy()))
            except TokenExpiredError:
                self.after(0, lambda: (self.destroy(), self.main_window._handle_session_expired()))
            except Exception as e:
                self.after(0, lambda: (
                    self.save_btn.config(state=NORMAL, text="Save"),
                    messagebox.showerror("Save failed", str(e))
                ))

        threading.Thread(target=do_save, daemon=True).start()


# ─── Change Password Dialog ───────────────────────────────────────────────────

class ChangePasswordDialog(Toplevel):
    def __init__(self, parent, api):
        super().__init__(parent)
        self.main_window = parent
        self.api = api
        self.title("Change My Password")
        self.geometry("360x230")
        self.resizable(False, False)
        self.grab_set()
        set_window_icon(self)

        frm = Frame(self, padx=24, pady=20)
        frm.pack(fill=BOTH, expand=True)

        Label(frm, text="Current Password").grid(row=0, column=0, sticky=W, pady=8)
        self.cur_pw = Entry(frm, show="*", width=26)
        self.cur_pw.grid(row=0, column=1, pady=8, padx=8)

        Label(frm, text="New Password").grid(row=1, column=0, sticky=W, pady=8)
        self.new_pw = Entry(frm, show="*", width=26)
        self.new_pw.grid(row=1, column=1, pady=8, padx=8)

        Label(frm, text="Confirm New").grid(row=2, column=0, sticky=W, pady=8)
        self.conf_pw = Entry(frm, show="*", width=26)
        self.conf_pw.grid(row=2, column=1, pady=8, padx=8)

        self.status_lbl = Label(frm, text="", fg="#dc2626", font=("Segoe UI", 9))
        self.status_lbl.grid(row=3, column=0, columnspan=2, pady=(4, 0))

        self.save_btn = Button(frm, text="Save Password", bg=PRIMARY_COLOR, fg="white", command=self._do_change)
        self.save_btn.grid(row=4, column=0, columnspan=2, sticky=EW, pady=10)

        self.cur_pw.focus_set()
        self.bind("<Return>", lambda e: self._do_change())

    def _do_change(self):
        if self.new_pw.get() != self.conf_pw.get():
            self.status_lbl.config(text="New passwords do not match.")
            return
        if len(self.new_pw.get()) < 8:
            self.status_lbl.config(text="New password must be at least 8 characters.")
            return
        self.save_btn.config(state=DISABLED, text="Saving...")
        cur, new = self.cur_pw.get(), self.new_pw.get()

        def attempt():
            try:
                self.api.change_own_password(cur, new)
                self.after(0, lambda: (messagebox.showinfo("Done", "Password changed successfully."), self.destroy()))
            except TokenExpiredError:
                self.after(0, lambda: (self.destroy(), self.main_window._handle_session_expired()))
            except Exception as e:
                self.after(0, lambda: (
                    self.save_btn.config(state=NORMAL, text="Save Password"),
                    self.status_lbl.config(text=str(e))
                ))

        threading.Thread(target=attempt, daemon=True).start()


class AdminResetPasswordDialog(Toplevel):
    """Owner/Business Manager resetting another user's forgotten password."""
    def __init__(self, parent, api, username):
        super().__init__(parent)
        self.main_window = parent
        self.api = api
        self.username = username
        self.title(f"Reset Password — {username}")
        self.geometry("360x230")
        self.resizable(False, False)
        self.grab_set()
        set_window_icon(self)

        frm = Frame(self, padx=24, pady=20)
        frm.pack(fill=BOTH, expand=True)

        Label(frm, text=f"New password for '{username}'", font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=2, pady=(0, 12))

        Label(frm, text="New Password").grid(row=1, column=0, sticky=W, pady=8)
        self.new_pw = Entry(frm, show="*", width=26)
        self.new_pw.grid(row=1, column=1, pady=8, padx=8)

        Label(frm, text="Confirm New").grid(row=2, column=0, sticky=W, pady=8)
        self.conf_pw = Entry(frm, show="*", width=26)
        self.conf_pw.grid(row=2, column=1, pady=8, padx=8)

        self.status_lbl = Label(frm, text="", fg="#dc2626", font=("Segoe UI", 9))
        self.status_lbl.grid(row=3, column=0, columnspan=2, pady=(4, 0))

        self.save_btn = Button(frm, text="Reset Password", bg=PRIMARY_COLOR, fg="white", command=self._do_reset)
        self.save_btn.grid(row=4, column=0, columnspan=2, sticky=EW, pady=10)

        self.new_pw.focus_set()
        self.bind("<Return>", lambda e: self._do_reset())

    def _do_reset(self):
        new, confirm = self.new_pw.get(), self.conf_pw.get()
        if new != confirm:
            self.status_lbl.config(text="Passwords do not match.")
            return
        if len(new) < 8:
            self.status_lbl.config(text="Password must be at least 8 characters.")
            return
        self.save_btn.config(state=DISABLED, text="Resetting...")

        def attempt():
            try:
                self.api.reset_user_password(self.username, new)
                self.after(0, lambda: (messagebox.showinfo("Done", f"Password reset for '{self.username}'."), self.destroy()))
            except TokenExpiredError:
                self.after(0, lambda: (self.destroy(), self.main_window._handle_session_expired()))
            except Exception as e:
                self.after(0, lambda: (
                    self.save_btn.config(state=NORMAL, text="Reset Password"),
                    self.status_lbl.config(text=str(e))
                ))

        threading.Thread(target=attempt, daemon=True).start()


# ─── Main Window ──────────────────────────────────────────────────────────────

class MainWindow(Tk):
    def __init__(self, api):
        super().__init__()
        self.api = api
        self.modules = api.modules()
        self.records = {}
        self.current_module = None
        self.title(f"{SOFTWARE_NAME} — {CLIENT_COMPANY_NAME} | {api.full_name} ({api.role})")
        self.configure(bg="#f8fafc")
        set_window_icon(self)

        # Restore saved window geometry
        prefs = load_prefs()
        self.geometry(prefs.get("geometry", "1300x760"))

        # Save geometry on close
        def _on_close():
            p = load_prefs()
            p["geometry"] = self.geometry()
            save_prefs(p)
            self.destroy()
        self.protocol("WM_DELETE_WINDOW", _on_close)

        self.header_app_logo = load_tk_image(APP_LOGO_PATH, (42, 42))
        self.header_company_logo = load_tk_image(COMPANY_LOGO_PATH, (120, 40))

        # ── Header bar ───────────────────────────────────────────────────────
        top = Frame(self, bg=PRIMARY_COLOR, height=60)
        top.pack(fill=X)
        if self.header_app_logo:
            Label(top, image=self.header_app_logo, bg=PRIMARY_COLOR).pack(side=LEFT, padx=(10, 6), pady=8)
        title_wrap = Frame(top, bg=PRIMARY_COLOR)
        title_wrap.pack(side=LEFT, padx=(4, 10), pady=8)
        Label(title_wrap, text=SOFTWARE_NAME, fg="white", bg=PRIMARY_COLOR, font=("Segoe UI", 16, "bold")).pack(anchor=W)
        Label(title_wrap, text=f"{CLIENT_COMPANY_NAME} | Developed by {DEVELOPER_NAME}", fg="#DBEAFE", bg=PRIMARY_COLOR, font=("Segoe UI", 9)).pack(anchor=W)
        if self.header_company_logo:
            Label(top, image=self.header_company_logo, bg=PRIMARY_COLOR).pack(side=LEFT, padx=10)

        # Right-side header buttons
        Button(top, text="Refresh", command=self.load_dashboard,
               bg=PRIMARY_COLOR, fg="white", relief=FLAT, padx=8).pack(side=RIGHT, padx=(4, 10), pady=12)
        Button(top, text="Change Password", command=self.open_change_password,
               bg="#0F3460", fg="white", relief=FLAT, padx=8).pack(side=RIGHT, padx=4, pady=12)

        # ── Notebook tabs ─────────────────────────────────────────────────────
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=BOTH, expand=True, padx=10, pady=10)
        self.dashboard_frame = Frame(self.notebook)
        self.notebook.add(self.dashboard_frame, text="Dashboard")

        self.tab_frames = {}
        for module in self.modules:
            frame = Frame(self.notebook)
            self.notebook.add(frame, text=module.replace("_", " ").title())
            self.tab_frames[module] = frame
            self.build_module_tab(frame, module)

        if self.api.role in ["Owner", "Business Manager"]:
            admin = Frame(self.notebook)
            self.notebook.add(admin, text="Admin")
            self.build_admin_tab(admin)

        help_frame = Frame(self.notebook)
        self.notebook.add(help_frame, text="Help & Guide")
        self.build_help_tab(help_frame)

        self.notebook.bind("<<NotebookTabChanged>>", self.on_tab_changed)
        self.load_dashboard()

    # ── Session expired ───────────────────────────────────────────────────────

    def _handle_session_expired(self):
        messagebox.showwarning(
            "Session Expired",
            "Your login session has expired (10-hour limit).\nPlease log in again."
        )
        self.destroy()
        LoginWindow(self.api).mainloop()

    # ── Change password ───────────────────────────────────────────────────────

    def open_change_password(self):
        ChangePasswordDialog(self, self.api)

    # ── Admin tab ─────────────────────────────────────────────────────────────

    def build_admin_tab(self, frame):
        nb = ttk.Notebook(frame)
        nb.pack(fill=BOTH, expand=True)

        users_frame = Frame(nb)
        audit_frame = Frame(nb)
        trash_frame = Frame(nb)
        create_frame = Frame(nb)
        nb.add(users_frame, text="Users")
        nb.add(audit_frame, text="Audit Log")
        nb.add(trash_frame, text="Trash")
        nb.add(create_frame, text="Create / Update User")

        self._build_users_subtab(users_frame)
        self._build_audit_subtab(audit_frame)
        self._build_trash_subtab(trash_frame)
        self._build_create_user_subtab(create_frame)

    def _build_users_subtab(self, frame):
        toolbar = Frame(frame)
        toolbar.pack(fill=X, pady=6)
        Button(toolbar, text="Refresh Users", command=self._load_user_list).pack(side=LEFT, padx=8)
        Button(toolbar, text="Reset Password", bg=PRIMARY_COLOR, fg="white", command=self._reset_selected_user_password).pack(side=LEFT, padx=4)
        Button(toolbar, text="Deactivate Selected", fg="#b91c1c", command=self._deactivate_user).pack(side=LEFT, padx=4)
        Button(toolbar, text="Activate Selected", fg="#166534", command=self._activate_user).pack(side=LEFT, padx=4)

        cols = ["username", "full_name", "role", "is_active", "created_at"]
        self.users_tree = ttk.Treeview(frame, columns=cols, show="headings", height=18)
        for col in cols:
            self.users_tree.heading(col, text=col.replace("_", " ").title())
            self.users_tree.column(col, width=160)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=self.users_tree.yview)
        self.users_tree.configure(yscrollcommand=vsb.set)
        self.users_tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)
        self._load_user_list()

    def _load_user_list(self):
        def fetch():
            try:
                users = self.api.list_users()
                self.after(0, lambda: self._populate_users(users))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=fetch, daemon=True).start()

    def _populate_users(self, users):
        for item in self.users_tree.get_children():
            self.users_tree.delete(item)
        for u in users:
            status = "Active" if u.get("is_active") else "Inactive"
            self.users_tree.insert("", END, iid=u["username"], values=[u["username"], u.get("full_name", ""), u.get("role", ""), status, u.get("created_at", "")])

    def _reset_selected_user_password(self):
        sel = self.users_tree.selection()
        if not sel:
            messagebox.showwarning("Select user", "Select a user to reset the password for.")
            return
        AdminResetPasswordDialog(self, self.api, sel[0])

    def _deactivate_user(self):
        sel = self.users_tree.selection()
        if not sel:
            messagebox.showwarning("Select user", "Select a user to deactivate.")
            return
        username = sel[0]
        if not messagebox.askyesno("Confirm", f"Deactivate user '{username}'?"):
            return
        def do():
            try:
                self.api.deactivate_user(username)
                self.after(0, lambda: (messagebox.showinfo("Done", f"'{username}' deactivated."), self._load_user_list()))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=do, daemon=True).start()

    def _activate_user(self):
        sel = self.users_tree.selection()
        if not sel:
            messagebox.showwarning("Select user", "Select a user to activate.")
            return
        username = sel[0]
        def do():
            try:
                self.api.activate_user(username)
                self.after(0, lambda: (messagebox.showinfo("Done", f"'{username}' activated."), self._load_user_list()))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=do, daemon=True).start()

    def _build_audit_subtab(self, frame):
        toolbar = Frame(frame)
        toolbar.pack(fill=X, pady=6)
        Label(toolbar, text="Module:").pack(side=LEFT, padx=(8, 2))
        self.audit_module_var = StringVar()
        mod_cb = ttk.Combobox(toolbar, textvariable=self.audit_module_var, values=[""] + list(self.modules.keys()), width=16)
        mod_cb.pack(side=LEFT, padx=4)
        Label(toolbar, text="User:").pack(side=LEFT, padx=(8, 2))
        self.audit_user_var = StringVar()
        Entry(toolbar, textvariable=self.audit_user_var, width=18).pack(side=LEFT, padx=4)
        Button(toolbar, text="Search", command=self._load_audit).pack(side=LEFT, padx=8)

        cols = ["id", "username", "action", "module", "record_id", "detail", "created_at"]
        self.audit_tree = ttk.Treeview(frame, columns=cols, show="headings", height=22)
        widths = {"id": 50, "username": 120, "action": 140, "module": 100, "record_id": 80, "detail": 160, "created_at": 160}
        for col in cols:
            self.audit_tree.heading(col, text=col.replace("_", " ").title())
            self.audit_tree.column(col, width=widths.get(col, 120))
        vsb = ttk.Scrollbar(frame, orient="vertical", command=self.audit_tree.yview)
        self.audit_tree.configure(yscrollcommand=vsb.set)
        self.audit_tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)

    def _load_audit(self):
        mod = self.audit_module_var.get().strip() or None
        uname = self.audit_user_var.get().strip() or None
        def fetch():
            try:
                rows = self.api.get_audit(module=mod, username=uname)
                self.after(0, lambda: self._populate_audit(rows))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=fetch, daemon=True).start()

    def _populate_audit(self, rows):
        for item in self.audit_tree.get_children():
            self.audit_tree.delete(item)
        for r in rows:
            self.audit_tree.insert("", END, values=[r.get("id"), r.get("username"), r.get("action"), r.get("module", ""), r.get("record_id", ""), r.get("detail", ""), r.get("created_at", "")])

    def _build_trash_subtab(self, frame):
        toolbar = Frame(frame)
        toolbar.pack(fill=X, pady=6)
        Label(toolbar, text="Module:").pack(side=LEFT, padx=(8, 2))
        self.trash_module_var = StringVar()
        mod_cb = ttk.Combobox(toolbar, textvariable=self.trash_module_var, values=list(self.modules.keys()), width=16)
        mod_cb.pack(side=LEFT, padx=4)
        if self.modules:
            mod_cb.set(list(self.modules.keys())[0])
        Button(toolbar, text="Load Trash", command=self._load_trash).pack(side=LEFT, padx=8)
        Button(toolbar, text="Restore Selected", fg="#166534", command=self._restore_record).pack(side=LEFT, padx=4)

        cols = ["id", "code", "title", "deleted_at"]
        self.trash_tree = ttk.Treeview(frame, columns=cols, show="headings", height=22)
        for col in cols:
            self.trash_tree.heading(col, text=col.replace("_", " ").title())
            self.trash_tree.column(col, width=200 if col != "id" else 60)
        vsb = ttk.Scrollbar(frame, orient="vertical", command=self.trash_tree.yview)
        self.trash_tree.configure(yscrollcommand=vsb.set)
        self.trash_tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)

    def _load_trash(self):
        module = self.trash_module_var.get().strip()
        if not module:
            messagebox.showwarning("Select module", "Choose a module.")
            return
        def fetch():
            try:
                recs = self.api.list_trash(module)
                self.after(0, lambda: self._populate_trash(recs))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=fetch, daemon=True).start()

    def _populate_trash(self, recs):
        for item in self.trash_tree.get_children():
            self.trash_tree.delete(item)
        for r in recs:
            self.trash_tree.insert("", END, iid=str(r["id"]), values=[r["id"], r.get("code", ""), r.get("title", ""), r.get("deleted_at", "")])

    def _restore_record(self):
        sel = self.trash_tree.selection()
        if not sel:
            messagebox.showwarning("Select record", "Select a record to restore.")
            return
        record_id = int(sel[0])
        def do():
            try:
                self.api.restore_record(record_id)
                self.after(0, lambda: (messagebox.showinfo("Restored", "Record restored."), self._load_trash()))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=do, daemon=True).start()

    def _build_create_user_subtab(self, frame):
        Label(frame, text="Create or Update User", font=("Segoe UI", 14, "bold")).pack(anchor=W, padx=16, pady=12)
        form = Frame(frame)
        form.pack(anchor=W, padx=16)
        labels = ["Username", "Full Name", "Role", "Password"]
        self.admin_entries = {}
        for i, lab in enumerate(labels):
            Label(form, text=lab).grid(row=i, column=0, sticky=W, pady=6)
            if lab == "Role":
                w = ttk.Combobox(form, values=["Owner", "Business Manager", "Sales Agent", "Ticketing Officer", "Finance Officer", "Viewer"], width=34)
            else:
                w = Entry(form, width=38, show="*" if lab == "Password" else "")
            w.grid(row=i, column=1, pady=6, padx=8)
            self.admin_entries[lab] = w
        Button(form, text="Save User", bg=PRIMARY_COLOR, fg="white", command=self.save_user).grid(row=5, column=1, sticky=EW, pady=12)
        Label(frame, text="Default first login: admin / admin123. Change immediately after setup.", fg="#b91c1c").pack(anchor=W, padx=16, pady=6)
        Label(frame, text="Set the licensed company name and logo in tenant_branding.json, then restart the client.", fg="#475569").pack(anchor=W, padx=16)

    def save_user(self):
        def do():
            try:
                self.api.create_user(
                    self.admin_entries["Username"].get(),
                    self.admin_entries["Full Name"].get(),
                    self.admin_entries["Role"].get(),
                    self.admin_entries["Password"].get(),
                )
                self.after(0, lambda: messagebox.showinfo("Saved", "User saved successfully."))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Failed", str(e)))
        threading.Thread(target=do, daemon=True).start()

    # ── Module tab ────────────────────────────────────────────────────────────

    def build_module_tab(self, frame, module):
        # Sort state
        frame.sort_col = None
        frame.sort_reverse = False

        # Toolbar
        toolbar = Frame(frame)
        toolbar.pack(fill=X, pady=6)
        Label(toolbar, text=f"{module.title()} Search:").pack(side=LEFT, padx=(8, 4))
        search = Entry(toolbar, width=36)
        search.pack(side=LEFT, padx=4)
        # Enter key triggers search
        search.bind("<Return>", lambda e, m=module, s=search: self.load_records(m, s.get()))
        Button(toolbar, text="Search", command=lambda m=module, s=search: self.load_records(m, s.get())).pack(side=LEFT, padx=4)
        Button(toolbar, text="Clear", command=lambda m=module, s=search: (s.delete(0, END), self.load_records(m))).pack(side=LEFT, padx=4)
        Button(toolbar, text="Add", bg=PRIMARY_COLOR, fg="white", command=lambda m=module: self.open_dialog(m)).pack(side=LEFT, padx=12)
        Button(toolbar, text="Edit Selected", command=lambda m=module: self.edit_selected(m)).pack(side=LEFT, padx=4)
        Button(toolbar, text="Delete Selected", command=lambda m=module: self.delete_selected(m)).pack(side=LEFT, padx=4)
        if module == "clients":
            Button(toolbar, text="New Quotation", command=self.new_quotation_for_client).pack(side=LEFT, padx=4)
        Button(toolbar, text="Export Excel", command=lambda m=module: self.export_excel(m)).pack(side=RIGHT, padx=8)
        if module == "clients":
            Button(toolbar, text="Import Excel", command=self.import_excel_clients).pack(side=RIGHT, padx=4)

        # Treeview
        cols = ["ID"] + self.modules[module]
        tree = ttk.Treeview(frame, columns=cols, show="headings")
        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        # Configure color tags
        for tag, (bg, fg) in ROW_TAGS.items():
            tree.tag_configure(tag, background=bg, foreground=fg)

        # Sortable column headings
        def sort_by(col, frm=frame, tr=tree):
            reverse = (frm.sort_col == col) and (not frm.sort_reverse)
            frm.sort_col = col
            frm.sort_reverse = reverse
            items = [(tr.set(item, col), item) for item in tr.get_children("")]
            try:
                items.sort(key=lambda t: float(t[0]) if t[0] else -float("inf"), reverse=reverse)
            except ValueError:
                items.sort(key=lambda t: t[0].lower(), reverse=reverse)
            for idx, (_, item) in enumerate(items):
                tr.move(item, "", idx)
            # Re-stripe after sort
            for i, item in enumerate(tr.get_children("")):
                if not tr.item(item, "tags"):
                    tr.item(item, tags=("oddrow" if i % 2 else "",))

        for col in cols:
            tree.heading(col, text=col, command=lambda c=col: sort_by(c))
            tree.column(col, width=140 if col != "ID" else 60, anchor=W)

        tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)
        hsb.pack(side=BOTTOM, fill=X)

        # Double-click to edit
        tree.bind("<Double-1>", lambda e, m=module: self.edit_selected(m))

        # Status bar showing record count
        status_bar = Label(frame, text="", anchor=W, fg="#475569", font=("Segoe UI", 9), padx=8)
        status_bar.pack(fill=X, side=BOTTOM)

        frame.tree = tree
        frame.search = search
        frame.status_bar = status_bar
        self.load_records(module)

    def on_tab_changed(self, _):
        text = self.notebook.tab(self.notebook.select(), "text").lower().replace(" ", "_")
        self.current_module = text if text in self.modules else None

    def refresh_current_tab(self):
        if self.current_module:
            self.load_records(self.current_module)
        self.load_dashboard()

    def load_records(self, module, q=""):
        frame = self.tab_frames[module]
        tree = frame.tree

        def fetch():
            try:
                recs = self.api.list_records(module, q)
                self.records[module] = recs
                self.after(0, lambda: self._populate_tree(tree, module, recs, frame))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Load failed", str(e)))

        threading.Thread(target=fetch, daemon=True).start()

    def _populate_tree(self, tree, module, recs, frame):
        for item in tree.get_children():
            tree.delete(item)
        fields = self.modules[module]
        for rec in recs:
            d = rec.get("data", {})
            values = [rec["id"]] + [d.get(f, "") for f in fields]
            tag = _row_tag(module, d)
            tree.insert("", END, iid=str(rec["id"]), values=values, tags=(tag,) if tag else ())

        # Update status bar
        count = len(recs)
        frame.status_bar.config(text=f"  {count} record{'s' if count != 1 else ''}")

        # Re-apply sort if one is active
        if frame.sort_col:
            col = frame.sort_col
            reverse = frame.sort_reverse
            items = [(tree.set(item, col), item) for item in tree.get_children("")]
            try:
                items.sort(key=lambda t: float(t[0]) if t[0] else -float("inf"), reverse=reverse)
            except ValueError:
                items.sort(key=lambda t: t[0].lower(), reverse=reverse)
            for idx, (_, item) in enumerate(items):
                tree.move(item, "", idx)

    def open_dialog(self, module, record=None):
        self.current_module = module
        RecordDialog(self, module, self.modules[module], self.api, record)

    def selected_record(self, module):
        tree = self.tab_frames[module].tree
        sel = tree.selection()
        if not sel:
            return None
        rid = int(sel[0])
        return next((r for r in self.records.get(module, []) if r["id"] == rid), None)

    def edit_selected(self, module):
        rec = self.selected_record(module)
        if not rec:
            messagebox.showwarning("Select record", "Please select a record first.")
            return
        self.open_dialog(module, rec)

    def delete_selected(self, module):
        rec = self.selected_record(module)
        if not rec:
            messagebox.showwarning("Select record", "Please select a record first.")
            return
        if not messagebox.askyesno("Confirm delete", "Delete selected record? Only Owner/Business Manager can delete."):
            return

        def do():
            try:
                self.api.delete_record(rec["id"])
                self.after(0, lambda: (self.load_records(module), self.load_dashboard()))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Delete failed", str(e)))

        threading.Thread(target=do, daemon=True).start()

    def export_excel(self, module):
        fields = self.modules[module]
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"{module}_export.xlsx",
        )
        if not path:
            return

        def do():
            try:
                # Always export all records regardless of active search filter
                recs = self.api.list_records(module, q="")
                wb = Workbook()
                ws = wb.active
                ws.title = module
                ws.append(["ID"] + fields)
                for r in recs:
                    ws.append([r["id"]] + [r.get("data", {}).get(f, "") for f in fields])
                wb.save(path)
                self.after(0, lambda: messagebox.showinfo("Exported", f"Saved {len(recs)} records to {path}"))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Export failed", str(e)))

        threading.Thread(target=do, daemon=True).start()

    def new_quotation_for_client(self):
        """Open a new Quotation prefilled with an existing client's name — no Lead required."""
        rec = self.selected_record("clients")
        if not rec:
            messagebox.showwarning("Select client", "Select a client first.")
            return
        name = rec.get("data", {}).get("Client Name", "")
        self.current_module = "quotations"
        RecordDialog(self, "quotations", self.modules["quotations"], self.api, record=None, prefill={"Client Name": name})

    def import_excel_clients(self):
        """Bulk-upload a company/client list from an .xlsx file into the Clients module."""
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not path:
            return
        fields = self.modules["clients"]

        def do():
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                header = next(rows, None)
                if not header:
                    raise ValueError("The spreadsheet is empty.")

                field_by_lower = {f.lower(): f for f in fields}
                col_map = {}
                for idx, h in enumerate(header):
                    if h is None:
                        continue
                    field = field_by_lower.get(str(h).strip().lower())
                    if field:
                        col_map[idx] = field
                if not col_map:
                    raise ValueError("No matching columns found. Expected headers: " + ", ".join(fields))

                existing = self.records.get("clients", [])
                existing_keys = {
                    (r.get("data", {}).get("Client Name", "").strip().lower(), r.get("data", {}).get("Company", "").strip().lower())
                    for r in existing
                }

                imported = skipped = failed = 0
                for row in rows:
                    if row is None or all(v is None for v in row):
                        continue
                    data = {f: "" for f in fields}
                    for idx, field in col_map.items():
                        if idx < len(row) and row[idx] is not None:
                            data[field] = str(row[idx]).strip()
                    name = data.get("Client Name", "").strip()
                    if not name:
                        continue
                    key = (name.lower(), data.get("Company", "").strip().lower())
                    if key in existing_keys:
                        skipped += 1
                        continue
                    try:
                        self.api.save_record("clients", fields, data)
                        existing_keys.add(key)
                        imported += 1
                    except TokenExpiredError:
                        raise
                    except Exception:
                        failed += 1

                msg = f"Imported {imported} client(s). Skipped {skipped} duplicate(s)."
                if failed:
                    msg += f" {failed} row(s) failed."
                self.after(0, lambda: (messagebox.showinfo("Import complete", msg), self.load_records("clients"), self.load_dashboard()))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Import failed", str(e)))

        threading.Thread(target=do, daemon=True).start()

    # ── Dashboard ─────────────────────────────────────────────────────────────

    def _jump_to_tab(self, module):
        """Switch notebook to the given module tab and refresh its records."""
        for tab_id in self.notebook.tabs():
            tab_text = self.notebook.tab(tab_id, "text").lower().replace(" ", "_")
            if tab_text == module:
                self.notebook.select(tab_id)
                self.current_module = module
                self.load_records(module)
                return

    def load_dashboard(self):
        for w in self.dashboard_frame.winfo_children():
            w.destroy()
        Label(self.dashboard_frame, text="Loading dashboard...", fg="#475569").pack(padx=20, pady=20)

        def fetch():
            try:
                data = self.api.dashboard()
                try:
                    reminders = self.api.reminders()
                except Exception:
                    reminders = {}
                self.after(0, lambda: self._render_dashboard(data, reminders))
            except TokenExpiredError:
                self.after(0, self._handle_session_expired)
            except Exception as e:
                self.after(0, lambda: self._dashboard_error(str(e)))

        threading.Thread(target=fetch, daemon=True).start()

    def _dashboard_error(self, msg):
        for w in self.dashboard_frame.winfo_children():
            w.destroy()
        Label(self.dashboard_frame, text=f"Dashboard load failed: {msg}", fg="red").pack(padx=20, pady=20)

    def _render_dashboard(self, data, reminders=None):
        for w in self.dashboard_frame.winfo_children():
            w.destroy()
        Label(self.dashboard_frame, text=f"{SOFTWARE_NAME} — Management Dashboard",
              font=("Segoe UI", 18, "bold"), fg=PRIMARY_COLOR).pack(anchor=W, padx=18, pady=16)

        # ── Metric cards (clickable) ──────────────────────────────────────────
        cards_frame = Frame(self.dashboard_frame)
        cards_frame.pack(fill=X, padx=18)

        metrics = [
            ("Open Leads",      data.get("open_leads"),      "#FEF3C7", "#92400E"),
            ("Overdue AR",      data.get("overdue_ar"),      "#FEE2E2", "#991B1B"),
            ("Open Complaints", data.get("open_complaints"), "#EDE9FE", "#5B21B6"),
            ("Open Exceptions", data.get("open_exceptions"), "#DBEAFE", "#1E40AF"),
        ]

        for label, value, bg, fg_color in metrics:
            module = CARD_MODULE_MAP.get(label)
            clickable = module and module in self.tab_frames

            card = Frame(cards_frame, bd=1, relief=SOLID, bg=bg, padx=18, pady=12)
            card.pack(side=LEFT, padx=8, pady=8, fill=X, expand=True)

            sub = Label(card, text=label, bg=bg, fg="#475569", font=("Segoe UI", 10, "bold"))
            sub.pack(anchor=W)
            val_lbl = Label(card, text=str(value), bg=bg, fg=fg_color, font=("Segoe UI", 28, "bold"))
            val_lbl.pack(anchor=W)

            if clickable:
                hint = Label(card, text="click to view →", bg=bg, fg="#94A3B8", font=("Segoe UI", 8))
                hint.pack(anchor=W)
                for widget in (card, sub, val_lbl, hint):
                    widget.configure(cursor="hand2")
                    widget.bind("<Button-1>", lambda e, m=module: self._jump_to_tab(m))

        # ── Record counts + Reminders side-by-side ───────────────────────────
        lower_frame = Frame(self.dashboard_frame)
        lower_frame.pack(fill=BOTH, expand=True, padx=18, pady=8)

        table = ttk.Treeview(lower_frame, columns=["Module", "Records"], show="headings", height=10)
        table.heading("Module", text="Module")
        table.heading("Records", text="Records")
        table.column("Module", width=200)
        table.column("Records", width=100)
        table.pack(side=LEFT, anchor=NW, padx=(0, 24))
        for module, count in data.get("counts", {}).items():
            table.insert("", END, values=[module.title(), count])

        if reminders:
            rem_frame = Frame(lower_frame)
            rem_frame.pack(side=LEFT, anchor=NW, fill=BOTH, expand=True)
            Label(rem_frame, text="Reminders & Alerts", font=("Segoe UI", 12, "bold"), fg="#b91c1c").pack(anchor=W)

            overdue_leads = reminders.get("overdue_leads", [])
            upcoming_leads = reminders.get("upcoming_leads", [])
            overdue_ar = reminders.get("overdue_ar", [])

            def _rem_section(parent, title, items, fg):
                if not items:
                    return
                Label(parent, text=title, font=("Segoe UI", 10, "bold"), fg=fg).pack(anchor=W, pady=(8, 2))
                for item in items[:10]:
                    text = f"  [{item.get('code','')}] {item.get('title','')} — {item.get('next_action_date') or item.get('due_date','')}"
                    Label(parent, text=text, font=("Segoe UI", 9), fg="#374151", wraplength=380, justify=LEFT).pack(anchor=W)

            _rem_section(rem_frame, f"Overdue Follow-Ups ({len(overdue_leads)})", overdue_leads, "#b91c1c")
            _rem_section(rem_frame, f"Upcoming Follow-Ups this week ({len(upcoming_leads)})", upcoming_leads, "#92400E")
            _rem_section(rem_frame, f"Overdue AR ({len(overdue_ar)})", overdue_ar, "#7C3AED")

            if not overdue_leads and not upcoming_leads and not overdue_ar:
                Label(rem_frame, text="No pending reminders.", fg="#16a34a", font=("Segoe UI", 10)).pack(anchor=W, pady=8)

    # ── Help tab ──────────────────────────────────────────────────────────────

    def build_help_tab(self, frame):
        prefs = load_prefs()
        lang_var = StringVar(value=prefs.get("help_lang", "en"))

        GUIDE = {
            "en": {
                "toggle_label": "ቋንቋ: አማርኛ",
                "sections": [
                    ("Getting Started", """
1. First Login
   • Default username: admin   Password: admin123
   • Change your password immediately from the Admin tab after first login.
   • Or use the "Change Password" button in the top-right at any time.
   • The client connects to the server over your office local network.
     Edit client_config.ini and set: url = http://<ServerPC_IP>:8080

2. Server Setup (on the main office PC)
   • Install MariaDB/MySQL and run sql/create_database.sql to create the database.
   • Run windows_scripts/01_install_server_dependencies.bat
   • Run windows_scripts/02_run_server.bat to start the server.
   • Test: open a browser and go to http://127.0.0.1:8080/health — you should see {"status":"ok"}

3. Client Setup (on each staff PC)
   • Edit client_config.ini — set the server IP address.
   • Run windows_scripts/03_install_client_dependencies.bat
   • Start the client: windows_scripts/04_run_client.bat

4. Branding Setup
   • Edit client/tenant_branding.json to set your company name and logo paths.
   • Restart the client after changes.
"""),
                    ("Module Guide", """
Dashboard
   • Shows live counts of open leads, overdue accounts, open complaints, and exceptions.
   • Click any metric card to jump directly to that module.
   • Click "Refresh" in the top-right to update all figures.

Leads
   • Track all incoming client inquiries. Assign Source, Priority, and Lead Status.
   • Rows are color-coded: red = Urgent/VIP priority, green = Closed-Won, grey = Closed-Lost.
   • Move leads to Quotations once they are ready to be priced.

Quotations
   • Create and track quotes. Set Quote Status from Draft → Sent → Accepted/Lost.
   • Rows are color-coded: green = Accepted, grey = Lost/Expired, yellow = Follow-Up.

Bookings
   • Record confirmed travel bookings. Track Issue Status and Payment Status.
   • Rows are color-coded: red = Overdue payment, green = Paid + Issued, grey = Cancelled.

Accounts Receivable (AR)
   • Track outstanding payments per client. Set AR Stage as payment progresses.
   • Overdue AR appears on the dashboard. Rows are color-coded: red = Overdue.

Complaints
   • Log client complaints with Severity and Status.
   • Rows are color-coded: red = Legal/Public severity, green = Resolved/Closed.

Exceptions
   • Record policy exceptions (discounts, refunds, credit approvals).
   • Rows are color-coded: green = Closed.

Clients
   • Master list of all client companies. Set Client Type (Retail, Corporate, VIP).

Suppliers
   • Manage airline, hotel, and tour operator contacts.

Admin (Owner / Business Manager only)
   • Create and manage staff user accounts.
   • Assign roles: Owner, Business Manager, Sales Agent, Ticketing Officer, Finance Officer, Viewer.
"""),
                    ("Search, Sort & Export", """
• Use the Search bar at the top of each module tab to filter records by keyword.
  Press Enter or click Search to apply the filter.
• Click Clear to show all records.
• Click any column header to sort by that column. Click again to reverse the sort.
• Double-click any row to open the edit dialog.
• The record count is shown at the bottom of each module tab.
• Click Export Excel to download a full export of all records in that module.
  (Export always exports ALL records regardless of the active search filter.)
"""),
                    ("Troubleshooting", """
Cannot reach server / Red status on login screen
   • Check that the server PC is powered on and the server is running.
   • Confirm the IP address in client_config.ini matches the server PC.
   • Check Windows Firewall: allow port 8080 inbound on the server PC.
   • Run: ping <ServerPC_IP> in a command prompt to verify network connectivity.

Login fails — "Too many attempts"
   • 5 failed logins in 5 minutes triggers a lockout. Wait 5 minutes and try again.
   • Ensure Caps Lock is off and you are using the correct username.

"Session Expired" message appears
   • Login sessions last 10 hours. After that, you will be prompted to log in again.
   • This is normal and expected — simply log in again to continue.

Records not saving / "Save failed" error
   • Check that the server is still running.
   • Check the server log for database errors.
   • Ensure your user account has the correct role for the action.

Export Excel produces an empty file
   • Make sure there are records in that module.
   • If a search filter is active when you click Export, it still exports all records.

Application is slow
   • If the network is congested, searches and loads may take a few seconds.
   • Close and restart the client if it becomes unresponsive.

Database backup
   • Run windows_scripts/06_mysql_backup.bat on the server PC.
   • Backups are saved to C:\\SA_US_ERP\\Backups.

Forgot admin password
   • Connect to MySQL directly and run:
     UPDATE sa_us_erp.users SET password_hash = '<new_hash>' WHERE username = 'admin';
   • Or drop and recreate the database using sql/create_database.sql (data loss).

Changing secret key / improving security
   • Edit server/server_config.ini and set a strong random value for secret_key.
   • Restart the server after changes.
   • Restrict network access to the office LAN only (do not expose port 8080 to the internet).
"""),
                    ("Security Notes", """
• Change the default admin password on first login (Admin tab or Change Password button).
• Use a strong secret_key in server_config.ini (at least 32 random characters).
• Change the default database password in server_config.ini.
• Do not expose the server port (8080) to the internet — office LAN only.
• Back up the database regularly using the backup script.
• Role-based access: assign the least-privileged role for each staff member.
  Viewer: read-only.  Sales Agent / Ticketing Officer: create and edit.
  Finance Officer: AR and payments.  Business Manager / Owner: full access + admin.
• Sessions expire after 10 hours — this is a security feature, not a bug.
"""),
                ],
            },
            "am": {
                "toggle_label": "Language: English",
                "sections": [
                    ("መጀመሪያ / Getting Started", """
1. ለመጀመሪያ ጊዜ መግባት
   • ነባሪ የተጠቃሚ ስም: admin   የይለፍ ቃል: admin123
   • ወዲያውኑ ከ Admin ትር ወይም "Change Password" ቁልፍ ሳጥን ላይ የይለፍ ቃልዎን ይቀይሩ።
   • ደንበኛው ከቢሮ አካባቢያዊ አውታረ መረብ ጋር ይገናኛል።
     client_config.ini ን ያርትዑ እና ያዘጋጁ: url = http://<ServerPC_IP>:8080

2. የሰርቨር ማዋቀር (በዋናው ቢሮ ፒሲ ላይ)
   • MariaDB/MySQL ን ይጫኑ እና sql/create_database.sql ን ያሂዱ።
   • windows_scripts/01_install_server_dependencies.bat ያሂዱ
   • windows_scripts/02_run_server.bat ን ያሂዱ።
   • ፈትሽ: http://127.0.0.1:8080/health — {"status":"ok"} ማየት አለቦት

3. የደንበኛ ማዋቀር (በእያንዳንዱ የሰራተኛ ፒሲ ላይ)
   • client_config.ini ን ያርትዑ — የሰርቨሩን IP አድራሻ ያዘጋጁ።
   • windows_scripts/03_install_client_dependencies.bat ያሂዱ
   • ደንበኛ ይጀምሩ: windows_scripts/04_run_client.bat

4. የምልክት ማዋቀር
   • client/tenant_branding.json ን ያርትዑ — የድርጅት ስምዎን እና አርማዎን ያዘጋጁ።
   • ለውጦቹ እንዲሰሩ ደንበኛ ይዝጉ እና እንደገና ይክፈቱ።
"""),
                    ("የሞጁል መመሪያ / Module Guide", """
ዳሽቦርድ (Dashboard)
   • ክፍት ሊዶችን፣ ዘግይቶ የቀረ ክፍያ፣ ክፍት ቅሬታዎችን እና ልዩ ፈቃዶችን ያሳያል።
   • ማናቸውንም ካርድ ጠቅ ያድርጉ ወደ해당 ሞጁል ለመሄድ።
   • "Refresh" ን ጠቅ ያድርጉ ዝማኔ ለማግኘት።

ሊዶች (Leads)
   • ሁሉንም ገቢ የደንበኛ ጥያቄዎች ይከታተሉ።
   • ቀለም ምልክት: ቀይ = ጣና/VIP ቅድሚያ፣ አረንጓዴ = Closed-Won፣ ግራጫ = Closed-Lost

ኮቴሽን (Quotations)
   • ዋጋ ፍለጋዎችን ይፍጠሩ እና ይከታተሉ። ሁኔታ: Draft → Sent → Accepted/Lost

ቡኪንግ (Bookings)
   • የተረጋገጡ የጉዞ ቡኪንጎችን ይዝግቡ።

የክፍያ ስርዓት (Accounts Receivable)
   • ያልተከፈሉ ክፍያዎችን በደንበኛ ይከታተሉ። ዘግይቶ የቀረ ክፍያ በዳሽቦርዱ ላይ ይታያል።

ቅሬታ (Complaints)
   • የደንበኛ ቅሬታዎችን ይዝግቡ። ደረጃና ሁኔታ ይምደቡ።

ልዩ ፈቃዶች (Exceptions)
   • ዋጋ ቅናሽ፣ ገንዘብ ተመላሽ፣ ብድር እና ሌሎች ፈቃዶችን ይዝግቡ።

ደንበኞች (Clients)
   • የሁሉም ደንበኛ ድርጅቶች ዋና ዝርዝር።

አቅራቢዎች (Suppliers)
   • የአየር መንገድ፣ ሆቴል እና የጉዞ አስደራጆች አድራሻ።

አስተዳደር (Admin) — ለ Owner / Business Manager ብቻ
   • የሰራተኛ መለያዎችን ይፍጠሩ እና ያስተዳድሩ።
"""),
                    ("ፍለጋ፣ ደረጃ እና ወደ ኤክሰል መላክ / Search, Sort & Export", """
• በእያንዳንዱ ሞጁል ትር አናት ላይ ባለው የፍለጋ ሳጥን ይፈልጉ።
  Enter ን ጫኑ ወይም Search ን ጠቅ ያድርጉ።
• Clear ን ጠቅ ያድርጉ ሁሉንም ዝርዝሮች ለማሳየት።
• ማናቸውንም የዓምድ ርዕስ ጠቅ ያድርጉ ለመደርደር። እንደገና ጠቅ ያድርጉ ለመገልበጥ።
• ማናቸውንም ረድፍ ሁለቴ ጠቅ ያድርጉ ለማርትዕ።
• Export Excel ን ጠቅ ያድርጉ ሁሉንም ዝርዝሮች ወደ ኤክሰል ፋይል ለመቅዳት።
"""),
                    ("ችግር ፈቺ / Troubleshooting", """
ሰርቨሩን ማግኘት አልቻልኩም / ቀይ ሁኔታ በግቤት ማያ ገጽ ላይ
   • የሰርቨሩ ፒሲ እንደ ተፃፈ እና ሰርቨሩ እየሄደ እንዳለ ያረጋግጡ።
   • client_config.ini ውስጥ ያለው IP አድራሻ ከሰርቨሩ ፒሲ ጋር ይዛመዳል።
   • Windows Firewall: ወደብ 8080 ን ፍቀዱ።

"Session Expired" መልዕክት
   • የግቤት ክፍለ ጊዜ 10 ሰዓት ነው። ካለፈ፣ እንደገና ይግቡ።

ዳታቤዝ ምትኬ (Backup)
   • windows_scripts/06_mysql_backup.bat ን ያሂዱ።
   • ምትኬዎች ወደ C:\\SA_US_ERP\\Backups ይቀመጣሉ።

የአስተዳዳሪ የይለፍ ቃል ረሳሁ
   • ቀጥታ ወደ MySQL ይገናኙ እና ያሂዱ:
     UPDATE sa_us_erp.users SET password_hash='<new_hash>' WHERE username='admin';
   • ወይም ዳታቤዙን ሰርዘው sql/create_database.sql ን ያሂዱ (ዳታ ይጠፋል)።
"""),
                    ("የደህንነት ማስታወሻዎች / Security Notes", """
• ለመጀመሪያ ጊዜ ሲገቡ ነባሪ የይለፍ ቃሉን ይቀይሩ (Admin ትር ወይም "Change Password" ቁልፍ)።
• server_config.ini ውስጥ ጠንካራ secret_key ይጠቀሙ (ቢያንስ 32 ዘፈቀደ ቁምፊዎች)።
• ነባሪ የዳታቤዝ የይለፍ ቃሉን ይቀይሩ።
• ወደብ 8080 ን ከኢንተርኔት አያጋልጡ — በቢሮ LAN ብቻ ያቆዩ።
• ዳታቤዙን በየጊዜው ምትኬ ያስቀምጡ።
• ክፍለ ጊዜ ከ10 ሰዓት በኋላ ያልቃል — ይህ የደህንነት ባህሪ ነው።
"""),
                ],
            },
        }

        def _render(lang):
            for w in content_frame.winfo_children():
                w.destroy()
            guide = GUIDE[lang]
            toggle_btn.config(text=guide["toggle_label"])
            for title, body in guide["sections"]:
                hdr = Frame(content_frame, bg="#EFF6FF")
                hdr.pack(fill=X, pady=(10, 0), padx=4)
                Label(hdr, text=title, font=("Segoe UI", 11, "bold"), fg=PRIMARY_COLOR,
                      bg="#EFF6FF", anchor=W).pack(fill=X, padx=8, pady=4)
                txt = Text(content_frame, wrap=WORD, font=("Nyala" if lang == "am" else "Segoe UI", 10),
                           relief=FLAT, bg="#f8fafc", height=len(body.strip().splitlines()) + 1,
                           state=NORMAL, padx=10, pady=4)
                txt.insert("1.0", body.strip())
                txt.config(state=DISABLED)
                txt.pack(fill=X, padx=12, pady=(0, 2))

        def _toggle():
            new_lang = "am" if lang_var.get() == "en" else "en"
            lang_var.set(new_lang)
            prefs = load_prefs()
            prefs["help_lang"] = new_lang
            save_prefs(prefs)
            _render(new_lang)

        top_bar = Frame(frame, bg="#f8fafc")
        top_bar.pack(fill=X, padx=10, pady=(8, 0))
        Label(top_bar, text="Help & Troubleshooting Guide", font=("Segoe UI", 14, "bold"),
              fg=PRIMARY_COLOR, bg="#f8fafc").pack(side=LEFT)
        toggle_btn = Button(top_bar, text="", command=_toggle,
                            bg=PRIMARY_COLOR, fg="white", relief=FLAT, padx=10, pady=3)
        toggle_btn.pack(side=RIGHT)

        canvas = Canvas(frame, bg="#f8fafc", highlightthickness=0)
        vscroll = ttk.Scrollbar(frame, orient="vertical", command=canvas.yview)
        content_frame = Frame(canvas, bg="#f8fafc")
        content_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=content_frame, anchor="nw")
        canvas.configure(yscrollcommand=vscroll.set)
        canvas.pack(side=LEFT, fill=BOTH, expand=True, padx=4, pady=4)
        vscroll.pack(side=RIGHT, fill=Y)
        canvas.bind("<MouseWheel>", lambda e: canvas.yview_scroll(int(-1 * (e.delta / 120)), "units"))

        _render(lang_var.get())


if __name__ == "__main__":
    api = API()
    LoginWindow(api).mainloop()
